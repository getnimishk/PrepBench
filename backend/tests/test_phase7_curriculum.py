# PrepBench - Copyright (c) 2026 Nimish Kanungo
# Licensed under the PolyForm Noncommercial License 1.0.0 (see LICENSE).
# Commercial use requires a separate licence from the copyright holder.

"""
test_phase7_curriculum.py

Phase 7, the Curriculum & Knowledge Layer (docs/implementation/PHASE-7-CONTRACT.md):

  * D1 -- the importer never takes a topic-number column for the title, and a roadmap
    already imported that way is repaired only on request: by phase and number, titles
    only, everything else kept, idempotent, refusing anything ambiguous.
  * D2 -- guide provenance is learner | ai | course. A course lesson stays one when
    edited; existing sections are relabelled only on request and only on an exact
    six-field match with a lesson file.
  * D3 / WP 7.6 -- roadmaps are scoped: "unassigned" lists only roadmaps with no
    preparation; a preparation's roadmap claim (roadmap_count) follows the real,
    unarchived link, so unlinking or archiving ends it.
  * D5 / WP 7.4 -- attached packs report what they hold, so a guide-only pack (ADLS)
    never claims scenarios.

Every test builds its own rows in the suite's throwaway database.
"""
import io
import json
import uuid
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.models.roadmap import RoadmapTopic, RoadmapTopicStatus
from tests.conftest import TestingSessionLocal

client = TestClient(app)

REPO = Path(__file__).resolve().parents[2]
ADF_WORKBOOK = REPO / "docs" / "research" / "ADF_Master_Roadmap_Mapped.xlsx"
LESSONS = REPO / "docs" / "research" / "agentic-ai" / "lessons"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook_rows():
    """(phase, number, title) for the ADF workbook's 60 topics, read independently of the app."""
    wb = openpyxl.load_workbook(ADF_WORKBOOK, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    header = next(i for i, r in enumerate(rows) if r and r[0] == "Phase")
    return [(str(r[0]).strip(), str(r[1]).strip(), r[2].strip()) for r in rows[header + 1:] if r and r[1] is not None]


def _xlsx(header, body) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for row in body:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _preparation(kind: str = "skill", pack: str | None = None) -> int:
    created = client.post("/api/v1/subjects", json={"name": f"P7 {uuid.uuid4().hex[:8]}", "kind": kind})
    assert created.status_code == 201, created.text
    sid = created.json()["id"]
    if pack:
        assert client.post(f"/api/v1/subjects/{sid}/content-packs", json={"pack_id": pack}).status_code == 201
    return sid


def _numbered_adf_roadmap(subject_id: int | None = None) -> dict:
    """ADF's roadmap as the old importer left it: every title the bare topic number."""
    topics = [{"title": number, "phase_name": phase} for phase, number, _ in _workbook_rows()]
    body = {"title": f"ADF Master Roadmap {uuid.uuid4().hex[:6]}", "topics": topics, "source_filename": "ADF_Master_Roadmap.xlsx"}
    made = client.post("/api/v1/roadmaps/import/confirm", json=body)
    assert made.status_code == 201, made.text
    rid = made.json()["roadmap_id"]
    if subject_id is not None:
        assert client.put(f"/api/v1/roadmaps/{rid}", json={"subject_id": subject_id}).status_code == 200
    return client.get(f"/api/v1/roadmaps/{rid}").json()


def _topics(detail: dict) -> list:
    return [t for p in detail["phases"] for t in p["topics"]]


def _repair(rid: int, step: str, content: bytes, name: str = "ADF_Master_Roadmap.xlsx", **data):
    return client.post(f"/api/v1/roadmaps/{rid}/title-repair/{step}", files={"file": (name, content, XLSX)}, data=data)


# ======================================================================= D1 -- importer


def test_the_adf_workbook_imports_with_its_real_topic_names():
    """The header is "Phase, Topic #, Topic, ...": the number column is never the title."""
    preview = client.post(
        "/api/v1/roadmaps/import/validate",
        files={"file": ("ADF_Master_Roadmap.xlsx", ADF_WORKBOOK.read_bytes(), XLSX)},
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    titles = [t["title"] for t in body["topics"]]
    expected = _workbook_rows()
    assert len(titles) == 60
    assert titles == [title for _, _, title in expected]
    assert not any(t.isdigit() for t in titles)
    assert len(body["phases"]) == 12
    assert [(t["phase_name"], t["title"]) for t in body["topics"]] == [(p, title) for p, _, title in expected]


@pytest.mark.parametrize("number_header", ["Topic #", "#", "No.", "Topic number", "Topic No", "Roadmap #"])
def test_a_number_column_never_becomes_the_title(number_header):
    content = _xlsx(["Phase", number_header, "Topic", "Learning Objective"],
                    [["1. Basics", 1, "First thing", "Learn it"], ["1. Basics", 2, "Second thing", "Learn it"]])
    body = client.post("/api/v1/roadmaps/import/validate", files={"file": ("plan.xlsx", content, XLSX)}).json()
    assert [t["title"] for t in body["topics"]] == ["First thing", "Second thing"]


def test_a_plain_topic_workbook_is_unchanged():
    content = _xlsx(["Phase", "Topic", "Learning Objective", "Est. Hours"],
                    [["1. Basics", "First thing", "Learn it", 2]])
    body = client.post("/api/v1/roadmaps/import/validate", files={"file": ("plan.xlsx", content, XLSX)}).json()
    assert [(t["title"], t["estimated_hours"]) for t in body["topics"]] == [("First thing", 2.0)]


def test_a_workbook_with_only_a_number_column_imports_no_numbers_as_titles():
    """No topic-name column means no topics -- never "1", "2", "3" as names."""
    content = _xlsx(["Phase", "Topic #", "Learning Objective", "Success Criteria"],
                    [["1. Basics", 1, "Learn it", "Explain it"]])
    body = client.post("/api/v1/roadmaps/import/validate", files={"file": ("plan.xlsx", content, XLSX)}).json()
    assert body["topics"] == []


def test_two_title_columns_keep_the_first_as_before():
    content = _xlsx(["Phase", "Topic", "Title", "Learning Objective"], [["1. Basics", "Named topic", "Other", "Learn"]])
    body = client.post("/api/v1/roadmaps/import/validate", files={"file": ("plan.xlsx", content, XLSX)}).json()
    assert [t["title"] for t in body["topics"]] == ["Named topic"]


# ================================================================== D1 -- title repair


def _make_topic_one_in_progress(rid: int, topic_id: int) -> None:
    """The learner's real state on ADF topic 1: in progress, a note, and here also a
    demonstration and a guide section, all of which the repair must leave alone."""
    assert client.patch(f"/api/v1/roadmaps/{rid}/topics/{topic_id}",
                        json={"status": "in_progress", "evidence_notes": "My own note on ADF."}).status_code == 200
    demo = client.post(f"/api/v1/roadmaps/{rid}/topics/{topic_id}/demonstrations",
                       json={"response_text": "ADF moves and orchestrates data across services.", "self_grade": "partial"})
    assert demo.status_code in (200, 201), demo.text
    section = client.post(f"/api/v1/roadmaps/{rid}/topics/{topic_id}/guide/sections",
                          json={"title": "My notes", "body": "ADF is an orchestrator."})
    assert section.status_code in (200, 201), section.text


def test_the_repair_previews_then_renames_only_titles_and_keeps_everything_else():
    sid = _preparation(pack="adf")
    detail = _numbered_adf_roadmap(sid)
    rid = detail["id"]
    topics = _topics(detail)
    first = next(t for t in topics if t["title"] == "1")
    _make_topic_one_in_progress(rid, first["id"])
    before = {t["id"]: t for t in _topics(client.get(f"/api/v1/roadmaps/{rid}").json())}
    demos_before = client.get(f"/api/v1/roadmaps/{rid}/topics/{first['id']}/demonstrations").json()
    guide_before = client.get(f"/api/v1/roadmaps/{rid}/topics/{first['id']}/guide").json()["sections"]

    preview = _repair(rid, "preview", ADF_WORKBOOK.read_bytes())
    assert preview.status_code == 200, preview.text
    plan = preview.json()
    assert plan["can_apply"] is True and plan["problems"] == [] and plan["already_named"] == 0
    assert len(plan["changes"]) == 60
    names = {(p, n): title for p, n, title in _workbook_rows()}
    for change in plan["changes"]:
        assert change["old_title"] == change["number"]
        assert change["new_title"] == names[(change["phase"], change["number"])]
    # The preview changed nothing.
    assert {t["id"]: t["title"] for t in _topics(client.get(f"/api/v1/roadmaps/{rid}").json())} == \
        {i: t["title"] for i, t in before.items()}

    ids = ",".join(str(c["topic_id"]) for c in plan["changes"])
    applied = _repair(rid, "apply", ADF_WORKBOOK.read_bytes(), topic_ids=ids)
    assert applied.status_code == 200, applied.text
    assert len(applied.json()["repaired"]) == 60

    after = {t["id"]: t for t in _topics(client.get(f"/api/v1/roadmaps/{rid}").json())}
    assert set(after) == set(before)  # same ids: nothing created or deleted
    for tid, t in after.items():
        old = before[tid]
        assert t["title"] == names[(next(p["name"] for p in detail["phases"] if p["id"] == old["phase_id"]), old["title"])]
        for field in ("status", "progress_percentage", "started_at", "completed_at", "evidence_notes",
                      "learning_objective", "success_criteria", "estimated_hours", "phase_id", "order_index"):
            assert t[field] == old[field], (tid, field)
    one = after[first["id"]]
    assert one["title"] == "What Azure Data Factory Is"
    assert one["status"] == "in_progress" and one["started_at"] is not None
    assert one["evidence_notes"] == "My own note on ADF."
    assert client.get(f"/api/v1/roadmaps/{rid}/topics/{first['id']}/demonstrations").json() == demos_before
    assert [s["id"] for s in client.get(f"/api/v1/roadmaps/{rid}/topics/{first['id']}/guide").json()["sections"]] == \
        [s["id"] for s in guide_before]
    # The chapters still map, now by name.
    assert client.get(f"/api/v1/roadmaps/{rid}/topics/{first['id']}/guide").json()["mapped_chapters"]


def test_the_repair_is_idempotent():
    rid = _numbered_adf_roadmap()["id"]
    plan = _repair(rid, "preview", ADF_WORKBOOK.read_bytes()).json()
    ids = ",".join(str(c["topic_id"]) for c in plan["changes"])
    assert _repair(rid, "apply", ADF_WORKBOOK.read_bytes(), topic_ids=ids).status_code == 200

    again = _repair(rid, "preview", ADF_WORKBOOK.read_bytes()).json()
    assert again["changes"] == [] and again["already_named"] == 60 and again["can_apply"] is False
    refused = _repair(rid, "apply", ADF_WORKBOOK.read_bytes(), topic_ids=ids)
    assert refused.status_code == 400


def test_the_repair_never_touches_another_roadmap():
    mine = _numbered_adf_roadmap()["id"]
    other = _numbered_adf_roadmap()
    plan = _repair(mine, "preview", ADF_WORKBOOK.read_bytes()).json()
    assert {c["topic_id"] for c in plan["changes"]}.isdisjoint({t["id"] for t in _topics(other)})
    ids = ",".join(str(c["topic_id"]) for c in plan["changes"])
    assert _repair(mine, "apply", ADF_WORKBOOK.read_bytes(), topic_ids=ids).status_code == 200
    assert all(t["title"].isdigit() for t in _topics(client.get(f"/api/v1/roadmaps/{other['id']}").json()))


def test_the_repair_refuses_a_confirmation_that_is_not_the_previewed_plan():
    rid = _numbered_adf_roadmap()["id"]
    plan = _repair(rid, "preview", ADF_WORKBOOK.read_bytes()).json()
    partial = ",".join(str(c["topic_id"]) for c in plan["changes"][:5])
    refused = _repair(rid, "apply", ADF_WORKBOOK.read_bytes(), topic_ids=partial)
    assert refused.status_code == 400
    assert all(t["title"].isdigit() for t in _topics(client.get(f"/api/v1/roadmaps/{rid}").json()))


def test_the_repair_refuses_an_ambiguous_workbook():
    rid = _numbered_adf_roadmap()["id"]
    rows = _workbook_rows()
    phase, number, _ = rows[0]
    content = _xlsx(["Phase", "Topic #", "Topic", "Learning Objective"],
                    [[p, int(n), t, "x"] for p, n, t in rows] + [[phase, int(number), "A different name", "x"]])
    plan = _repair(rid, "preview", content).json()
    assert plan["can_apply"] is False
    assert any("more than one topic" in p for p in plan["problems"])


def test_the_repair_refuses_a_workbook_missing_a_topic():
    rid = _numbered_adf_roadmap()["id"]
    rows = _workbook_rows()[:-1]
    content = _xlsx(["Phase", "Topic #", "Topic", "Learning Objective"], [[p, int(n), t, "x"] for p, n, t in rows])
    plan = _repair(rid, "preview", content).json()
    assert plan["can_apply"] is False
    assert any("no row in the workbook" in p for p in plan["problems"])
    ids = ",".join(str(c["topic_id"]) for c in plan["changes"])
    assert _repair(rid, "apply", content, topic_ids=ids).status_code == 400


def test_the_repair_refuses_a_name_that_would_change_another_field():
    """An inconsistent row (in progress, never stamped as started) would be corrected by
    the write path; the repair refuses rather than change anything but a title."""
    rid = _numbered_adf_roadmap()["id"]
    db = TestingSessionLocal()
    try:
        topic = db.query(RoadmapTopic).filter(RoadmapTopic.roadmap_id == rid, RoadmapTopic.title == "5").first()
        topic.status = RoadmapTopicStatus.IN_PROGRESS
        topic.started_at = None
        db.commit()
    finally:
        db.close()
    plan = _repair(rid, "preview", ADF_WORKBOOK.read_bytes()).json()
    ids = ",".join(str(c["topic_id"]) for c in plan["changes"])
    refused = _repair(rid, "apply", ADF_WORKBOOK.read_bytes(), topic_ids=ids)
    assert refused.status_code == 400 and "started_at" in refused.json()["detail"]
    assert all(t["title"].isdigit() for t in _topics(client.get(f"/api/v1/roadmaps/{rid}").json()))


# ===================================================================== D2 -- provenance


def _lessons(n: int = 3) -> list:
    files = sorted(LESSONS.glob("*.guide.json"))[:n]
    return [json.loads(f.read_text(encoding="utf-8")) for f in files]


def _guide_roadmap_with_lessons(lessons: list) -> tuple:
    rid = client.post("/api/v1/roadmaps", json={"title": f"Agentic {uuid.uuid4().hex[:6]}"}).json()["id"]
    phase = client.post(f"/api/v1/roadmaps/{rid}/phases", json={"name": "Phase 1"}).json()["id"]
    ids = []
    for lesson in lessons:
        tid = client.post(f"/api/v1/roadmaps/{rid}/topics", json={"phase_id": phase, "title": lesson["topic_title"]}).json()["id"]
        for section in lesson["sections"]:
            # Loaded the old way: no source, so "learner" -- the mislabel D2 corrects.
            made = client.post(f"/api/v1/roadmaps/{rid}/topics/{tid}/guide/sections", json=section)
            assert made.status_code in (200, 201), made.text
            ids.append((tid, made.json()["id"]))
    return rid, ids


def test_a_course_section_is_created_as_course_and_stays_course_when_edited():
    rid, _ = _guide_roadmap_with_lessons([])
    phase = client.post(f"/api/v1/roadmaps/{rid}/phases", json={"name": "P"}).json()["id"]
    tid = client.post(f"/api/v1/roadmaps/{rid}/topics", json={"phase_id": phase, "title": "T"}).json()["id"]
    base = f"/api/v1/roadmaps/{rid}/topics/{tid}/guide/sections"
    made = client.post(base, json={"title": "Lesson", "body": "Course text.", "source": "course"}).json()
    assert made["source"] == "course" and made["edited_at"] is None

    same = client.put(f"{base}/{made['id']}", json={"title": "Lesson", "body": "Course text."}).json()
    assert same["source"] == "course" and same["edited_at"] is None  # saving unchanged is not an edit
    edited = client.put(f"{base}/{made['id']}", json={"title": "Lesson", "body": "Course text, my wording."}).json()
    assert edited["source"] == "course" and edited["edited_at"] is not None

    mine = client.post(base, json={"title": "Mine", "body": "My words."}).json()
    assert mine["source"] == "learner"
    assert client.post(base, json={"title": "X", "body": "Y", "source": "ai"}).status_code == 422


def test_relabel_proposes_only_exact_six_field_matches_and_changes_only_source():
    lessons = _lessons(3)
    rid, ids = _guide_roadmap_with_lessons(lessons)
    tid, changed_id = ids[0]
    # One section edited after loading: it no longer matches any lesson exactly.
    original = client.get(f"/api/v1/roadmaps/{rid}/topics/{tid}/guide").json()["sections"][0]
    client.put(f"/api/v1/roadmaps/{rid}/topics/{tid}/guide/sections/{changed_id}",
               json={**{k: original[k] for k in ("title", "body", "example", "common_mistake", "check_question", "check_answer")},
                     "body": original["body"] + " Edited."})

    preview = client.post(f"/api/v1/roadmaps/{rid}/guide/course-lessons/preview", json={"lessons": lessons})
    assert preview.status_code == 200, preview.text
    plan = preview.json()
    matched = {m["section_id"] for m in plan["matched"]}
    total = sum(len(lesson["sections"]) for lesson in lessons)
    assert changed_id not in matched
    assert len(matched) == total - 1 and plan["unmatched_written_by_you"] == 1

    before = {s["id"]: s for t, _ in ids for s in client.get(f"/api/v1/roadmaps/{rid}/topics/{t}/guide").json()["sections"]}
    applied = client.post(f"/api/v1/roadmaps/{rid}/guide/course-lessons/apply",
                          json={"lessons": lessons, "section_ids": sorted(matched)})
    assert applied.status_code == 200, applied.text
    assert len(applied.json()["relabelled"]) == total - 1
    after = {s["id"]: s for t, _ in ids for s in client.get(f"/api/v1/roadmaps/{rid}/topics/{t}/guide").json()["sections"]}
    for sid, s in after.items():
        assert s["source"] == ("learner" if sid == changed_id else "course")
        for field in ("title", "body", "example", "common_mistake", "check_question", "check_answer", "edited_at", "read_at"):
            assert s[field] == before[sid][field], (sid, field)


def test_relabel_refuses_a_section_that_is_not_an_exact_match():
    lessons = _lessons(1)
    rid, ids = _guide_roadmap_with_lessons(lessons)
    tid = ids[0][0]
    stray = client.post(f"/api/v1/roadmaps/{rid}/topics/{tid}/guide/sections", json={"title": "Mine", "body": "My words."}).json()
    refused = client.post(f"/api/v1/roadmaps/{rid}/guide/course-lessons/apply",
                          json={"lessons": lessons, "section_ids": [ids[0][1], stray["id"]]})
    assert refused.status_code == 400
    assert all(s["source"] == "learner" for s in client.get(f"/api/v1/roadmaps/{rid}/topics/{tid}/guide").json()["sections"])


def test_one_unmatched_section_can_be_confirmed_on_its_own_and_an_ai_draft_never_relabelled():
    rid, _ = _guide_roadmap_with_lessons([])
    phase = client.post(f"/api/v1/roadmaps/{rid}/phases", json={"name": "P"}).json()["id"]
    tid = client.post(f"/api/v1/roadmaps/{rid}/topics", json={"phase_id": phase, "title": "T"}).json()["id"]
    base = f"/api/v1/roadmaps/{rid}/topics/{tid}/guide/sections"
    section = client.post(base, json={"title": "Lesson", "body": "Course text, edited once."}).json()
    confirmed = client.put(f"{base}/{section['id']}",
                           json={"title": "Lesson", "body": "Course text, edited once.", "source": "course"}).json()
    assert confirmed["source"] == "course" and confirmed["edited_at"] is None

    db = TestingSessionLocal()
    try:
        from app.models.roadmap import TopicGuideSection
        draft = TopicGuideSection(topic_id=tid, order_index=9, title="Draft", body="AI text", source="ai")
        db.add(draft)
        db.commit()
        draft_id = draft.id
    finally:
        db.close()
    refused = client.put(f"{base}/{draft_id}", json={"title": "Draft", "body": "AI text", "source": "course"})
    assert refused.status_code == 400


# ================================================================ WP 7.6 / D3 -- scoping


def test_unassigned_lists_only_roadmaps_with_no_preparation():
    sid = _preparation()
    owned = client.post("/api/v1/roadmaps", json={"title": "Owned", "subject_id": sid}).json()["id"]
    loose = client.post("/api/v1/roadmaps", json={"title": "Loose"}).json()["id"]
    unassigned = client.get("/api/v1/roadmaps", params={"unassigned": True}).json()
    assert loose in {r["id"] for r in unassigned} and owned not in {r["id"] for r in unassigned}
    assert all(r["subject_id"] is None for r in unassigned)
    mine = client.get("/api/v1/roadmaps", params={"subject_id": sid}).json()
    assert {r["id"] for r in mine} == {owned}
    assert client.get("/api/v1/roadmaps", params={"subject_id": sid, "unassigned": True}).status_code == 400


def test_two_preparations_of_one_kind_see_only_their_own_roadmaps():
    a, b = _preparation(pack="adf"), _preparation(pack="adf")
    ra = client.post("/api/v1/roadmaps", json={"title": "A plan", "subject_id": a}).json()["id"]
    rb = client.post("/api/v1/roadmaps", json={"title": "B plan", "subject_id": b}).json()["id"]
    assert {r["id"] for r in client.get("/api/v1/roadmaps", params={"subject_id": a}).json()} == {ra}
    assert {r["id"] for r in client.get("/api/v1/roadmaps", params={"subject_id": b}).json()} == {rb}


def _subject(sid: int) -> dict:
    return client.get(f"/api/v1/subjects/{sid}").json()


def test_the_roadmap_claim_follows_the_real_link():
    """D3: unlink or archive a roadmap and the preparation stops claiming one; link it
    to another preparation and that one does."""
    psm, databricks = _preparation(), _preparation()
    rid = client.post("/api/v1/roadmaps", json={"title": "Storage to Cloud", "subject_id": psm}).json()["id"]
    assert _subject(psm)["roadmap_count"] == 1 and _subject(databricks)["roadmap_count"] == 0

    assert client.put(f"/api/v1/roadmaps/{rid}", json={"subject_id": None}).status_code == 200  # unlink
    assert _subject(psm)["roadmap_count"] == 0
    assert rid in {r["id"] for r in client.get("/api/v1/roadmaps", params={"unassigned": True}).json()}

    assert client.put(f"/api/v1/roadmaps/{rid}", json={"subject_id": databricks}).status_code == 200  # link
    assert _subject(databricks)["roadmap_count"] == 1

    assert client.put(f"/api/v1/roadmaps/{rid}", json={"is_archived": True}).status_code == 200  # archive
    assert _subject(databricks)["roadmap_count"] == 0


# ==================================================================== D5 / WP 7.4 -- packs


def test_attached_packs_report_what_they_hold():
    sid = _preparation(pack="adls")
    packs = {p["pack_id"]: p for p in _subject(sid)["content_packs"]}
    assert packs["adls"]["chapter_count"] == 11
    assert packs["adls"]["written_scenario_count"] == 0  # a guide, never a scenario source
    adf = _preparation(pack="adf")
    adf_pack = {p["pack_id"]: p for p in _subject(adf)["content_packs"]}["adf"]
    assert adf_pack["chapter_count"] == 21 and adf_pack["written_scenario_count"] > 0


def test_a_guide_only_pack_maps_no_roadmap_chapters():
    """ADLS has no roadmap alignment, so a preparation with only ADLS links no chapters."""
    sid = _preparation(pack="adls")
    rid = client.post("/api/v1/roadmaps", json={"title": "Plan", "subject_id": sid}).json()["id"]
    phase = client.post(f"/api/v1/roadmaps/{rid}/phases", json={"name": "P"}).json()["id"]
    tid = client.post(f"/api/v1/roadmaps/{rid}/topics", json={"phase_id": phase, "title": "Hierarchical Namespace"}).json()["id"]
    assert client.get(f"/api/v1/roadmaps/{rid}/topics/{tid}/guide").json()["mapped_chapters"] == []
    assert client.get(f"/api/v1/roadmaps/{rid}").json()["linked_pack_id"] is None


def test_a_linked_roadmap_names_the_pack_version_its_preparation_pins():
    """The topic page reads scenarios from this version, so it must be the pinned one."""
    sid = _preparation(pack="adf")
    pinned = {p["pack_id"]: p for p in _subject(sid)["content_packs"]}["adf"]["pack_version"]
    detail = _numbered_adf_roadmap(sid)
    assert (detail["linked_pack_id"], detail["linked_pack_version"]) == ("adf", pinned)
    listed = {r["id"]: r for r in client.get("/api/v1/roadmaps", params={"subject_id": sid}).json()}
    assert listed[detail["id"]]["linked_pack_version"] == pinned
    # Linked to no preparation: no pack, so no version either.
    loose = _numbered_adf_roadmap()
    assert (loose["linked_pack_id"], loose["linked_pack_version"]) == (None, None)


# ========================================================================= D4 -- unchanged


def test_completion_is_still_only_a_demonstration():
    rid = client.post("/api/v1/roadmaps", json={"title": f"D4 {uuid.uuid4().hex[:6]}"}).json()["id"]
    phase = client.post(f"/api/v1/roadmaps/{rid}/phases", json={"name": "P"}).json()["id"]
    tid = client.post(f"/api/v1/roadmaps/{rid}/topics", json={"phase_id": phase, "title": "T"}).json()["id"]
    assert client.patch(f"/api/v1/roadmaps/{rid}/topics/{tid}", json={"status": "completed"}).status_code == 400
    done = client.post(f"/api/v1/roadmaps/{rid}/topics/{tid}/demonstrations",
                       json={"response_text": "I can explain this topic without notes, fully.", "self_grade": "yes"})
    assert done.status_code in (200, 201) and done.json()["topic"]["status"] == "completed"


def test_workspace_counts_each_provenance_as_itself():
    sid = _preparation()
    rid = client.post("/api/v1/roadmaps", json={"title": "Plan", "subject_id": sid}).json()["id"]
    phase = client.post(f"/api/v1/roadmaps/{rid}/phases", json={"name": "P"}).json()["id"]
    tid = client.post(f"/api/v1/roadmaps/{rid}/topics", json={"phase_id": phase, "title": "T"}).json()["id"]
    base = f"/api/v1/roadmaps/{rid}/topics/{tid}/guide/sections"
    client.post(base, json={"title": "Mine", "body": "My words."})
    client.post(base, json={"title": "Lesson", "body": "Course text.", "source": "course"})
    db = TestingSessionLocal()
    try:
        from app.models.roadmap import TopicGuideSection
        db.add(TopicGuideSection(topic_id=tid, order_index=5, title="Draft", body="AI text", source="ai"))
        db.commit()
    finally:
        db.close()
    item = next(i for i in client.get("/api/v1/workspace", params={"subject_id": sid}).json()["items"]
                if i["kind"] == "topic_guide")
    assert item["detail"] == "3 sections: 1 written by you, 1 course lesson, 1 AI draft"
