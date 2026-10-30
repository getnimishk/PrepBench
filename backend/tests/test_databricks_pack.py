# PrepBench - Copyright (c) 2026 Nimish Kanungo
# Licensed under the PolyForm Noncommercial License 1.0.0 (see LICENSE).
# Commercial use requires a separate licence from the copyright holder.

"""
test_databricks_pack.py

The Databricks study guide (`backend/app/content/packs/databricks/`) is written as
one Markdown file per chapter under `src/` and built into the newest `v<N>.json` by
`scripts/build_content_pack.py`. These tests hold the two together and hold the
guide to its sources:

  * the shipped JSON is exactly what the sources build to;
  * every topic of the Databricks Data Engineering Mastery Roadmap
    (docs/research/Databricks_Data_Engineering_Mastery_Roadmap.xlsx, the learner
    imports it) maps to a chapter, by number and by its exact title;
  * every chapter cites sections that exist in
    docs/research/databricks-documentation-notes.md, where each fact is traced to an
    official Databricks, Microsoft, Delta Lake, Apache Parquet or Apache Spark page;
  * renamed products appear only as former names, Unity Catalog SQL is never said to
    have a DENY statement, and no exam pass mark is stated (Databricks publishes none);
  * v1 (the guide alone) stays as published, and v2's 18 written scenarios (one JSON
    file each under `src/scenarios/`) have the shape the scenario page reads, cite
    real notes sections, and hold to the same naming, DENY and pass-mark rules.
"""
import re
import sys
from pathlib import Path

import openpyxl

from app.content import packs as pack_store
from app.services.content_pack_service import get_pack_roadmap_alignments, normalize_topic_title

BACKEND = Path(__file__).resolve().parents[1]
REPO = BACKEND.parent
sys.path.insert(0, str(BACKEND / "scripts"))

import build_content_pack  # noqa: E402

NOTES = REPO / "docs" / "research" / "databricks-documentation-notes.md"
ROADMAP = REPO / "docs" / "research" / "Databricks_Data_Engineering_Mastery_Roadmap.xlsx"

# The version the sources build.
LATEST = 2

# Former product names (notes §0). They may appear only where the text says they are former names.
FORMER_NAMES = [
    "Delta Live Tables",
    "Databricks Asset Bundles",
    "Delta Sharing",
    "APPLY CHANGES",
    "Repos",
    "Delta cache",
    "DBIO cache",
    "Lakeflow Declarative Pipelines",
]
FORMER_MARKERS = re.compile(
    r"formerly|former|old name|older|replaced|renamed|is now|now called|still works|under a new name",
    re.IGNORECASE,
)


def _roadmap_topics():
    """(number, title) for every topic of the roadmap's syllabus sheet, in order."""
    sheet = openpyxl.load_workbook(ROADMAP, read_only=True).worksheets[0]
    return [(r[1], r[2]) for r in sheet.iter_rows(min_row=2, values_only=True) if isinstance(r[1], int)]


def _pack():
    pack = pack_store.get("databricks", LATEST)
    assert pack is not None, "the databricks pack failed to load"
    return pack


def _sentences(text: str):
    return re.split(r"(?<=[.!?])\s+|\n", text)


def test_the_shipped_json_is_what_the_sources_build_to():
    built = build_content_pack.render(build_content_pack.build("databricks"))
    shipped = (BACKEND / "app" / "content" / "packs" / "databricks" / f"v{LATEST}.json").read_text(encoding="utf-8")
    assert built == shipped, "the newest databricks pack is out of date: run scripts/build_content_pack.py databricks"


def test_the_roadmap_is_numbered_one_to_n():
    topics = _roadmap_topics()
    assert len(topics) == 108
    assert [n for n, _ in topics] == list(range(1, len(topics) + 1))


def test_every_roadmap_topic_maps_to_a_chapter_by_number_and_title():
    alignments = get_pack_roadmap_alignments("databricks", LATEST)
    topics = _roadmap_topics()
    for number, title in topics:
        assert alignments.get(normalize_topic_title(title)), f"Topic {number} {title!r} has no chapter"
        by_number = alignments[str(number)]
        assert {normalize_topic_title(m.topic_title) for m in by_number} == {normalize_topic_title(title)}, (
            f"Topic {number} is aligned under another title: {[m.topic_title for m in by_number]}"
        )
    numbers = {m.topic_number for ms in alignments.values() for m in ms}
    assert numbers == {n for n, _ in topics}


def test_coverage_is_stated_for_every_mapping():
    for mappings in get_pack_roadmap_alignments("databricks", LATEST).values():
        for m in mappings:
            assert m.coverage in {"Full", "Partial"}
            assert m.relevant_sections and m.learning_evidence


def test_learning_evidence_is_the_roadmap_success_criterion():
    """The chapter asks for the same proof as the roadmap the learner tracks."""
    sheet = openpyxl.load_workbook(ROADMAP, read_only=True).worksheets[0]
    criteria = {r[1]: r[4] for r in sheet.iter_rows(min_row=2, values_only=True) if isinstance(r[1], int)}
    alignments = get_pack_roadmap_alignments("databricks", LATEST)
    for number, criterion in criteria.items():
        assert {m.learning_evidence for m in alignments[str(number)]} == {criterion}, f"Topic {number}"


def test_each_chapter_opens_with_its_alignment_and_cites_real_notes_sections():
    notes = NOTES.read_text(encoding="utf-8")
    sections = {int(n) for n in re.findall(r"^## (\d+)\. ", notes, re.M)}
    for chapter in _pack().chapters:
        assert chapter.blocks[0].heading == "Roadmap alignment", chapter.id
        cited = {int(n) for n in re.findall(r"§(\d+)", chapter.sources)}
        assert cited, f"{chapter.id} cites no notes section"
        assert cited <= sections, f"{chapter.id} cites missing sections {sorted(cited - sections)}"


def test_relevant_sections_name_headings_that_exist():
    for chapter in _pack().chapters:
        headings = {b.heading for b in chapter.blocks if b.heading}
        for line in re.findall(r"\*Relevant sections:\*\s*(.+)", chapter.blocks[0].md or ""):
            for part in (p.strip() for p in line.split(";")):
                assert part in headings, f"{chapter.id}: no heading {part!r}"


def test_every_chapter_ends_with_a_trade_off_and_an_interview_takeaway():
    for chapter in _pack().chapters:
        headings = [b.heading for b in chapter.blocks]
        assert headings[-2:] == ["Design trade-off", "Interview takeaway"], chapter.id


def test_every_table_header_cell_has_text():
    """An empty header cell fails axe's empty-table-header rule on the chapter page."""
    for chapter in _pack().chapters:
        for block in chapter.blocks:
            if block.table:
                assert all(cell.strip() for cell in block.table.head), f"{chapter.id}/{block.heading}"


def test_every_figure_the_guide_shows_exists():
    public = REPO / "frontend" / "public"
    figures = []
    for chapter in _pack().chapters:
        for block in chapter.blocks:
            figures += re.findall(r"^!\[[^\]]+\]\((/assets/[^)\s]+)\)$", block.md or "", re.M)
    assert len(figures) == 4
    for path in figures:
        assert (public / path.lstrip("/")).is_file(), path


def test_former_product_names_appear_only_as_former_names():
    """Databricks renamed many products in 2025-26 (notes §0): the guide teaches the current names."""
    for chapter in _pack().chapters:
        for block in chapter.blocks:
            if block.heading in ("Roadmap alignment", "Names you will meet in older material"):
                continue
            for sentence in _sentences(block.md or ""):
                for name in FORMER_NAMES:
                    if re.search(r"\b" + re.escape(name) + r"\b", sentence):
                        assert FORMER_MARKERS.search(sentence), f"{chapter.id}: {name!r} in {sentence!r}"


def test_deny_is_never_taught_as_unity_catalog_sql():
    """Unity Catalog SQL has no DENY statement (notes §6.3, §19)."""
    for chapter in _pack().chapters:
        for block in chapter.blocks:
            if block.heading == "Roadmap alignment":
                continue  # names the roadmap topic "Ownership and the Missing DENY", nothing more
            for sentence in _sentences(block.md or ""):
                if "DENY" in sentence:
                    assert re.search(r"no `?DENY|legacy|ABAC|Beta|exam|DENY policy", sentence), (
                        f"{chapter.id}: {sentence!r}"
                    )


def test_no_exam_pass_mark_is_stated():
    """Databricks publishes no pass mark for its data engineering exams (notes §1, §19)."""
    for chapter in _pack().chapters:
        for block in chapter.blocks:
            assert not re.search(r"pass(ing)? (mark|score)", block.md or "", re.IGNORECASE), chapter.id


# Databricks v1, the guide without scenarios as first published (PR #116). A skill pinned at v1 keeps reading it until
# the learner upgrades, so the file must not change: publish a new version instead.
DATABRICKS_V1_FINGERPRINT = "4e33c8b5c6389c8a05fa6256bc1d89a0fc7b7c1befab1d3551fd954ecb082ce2"

ROLES = {"po", "pm", "dm", "em"}


def _canonical_hash(pack) -> str:
    import hashlib
    import json

    payload = json.dumps(pack.model_dump(mode="json"), sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _scenarios():
    return [s for level in _pack().scenario_levels for s in level.scenarios]


def _scenario_sentences(scenario):
    """Every sentence a learner reads in a scenario, from any string anywhere in it."""
    def strings(node):
        if isinstance(node, str):
            yield node
        elif isinstance(node, list):
            for item in node:
                yield from strings(item)
        elif isinstance(node, dict):
            for item in node.values():
                yield from strings(item)

    for text in strings([scenario.title, scenario.outcome, scenario.content]):
        yield from _sentences(text)


def test_the_first_published_databricks_guide_is_kept_unchanged():
    v1 = pack_store.get("databricks", 1)
    assert v1 is not None and v1.scenario_levels == []
    assert _canonical_hash(v1) == DATABRICKS_V1_FINGERPRINT
    # The newest version keeps every chapter, in order, so links and pinned progress still resolve.
    assert [c.id for c in v1.chapters] == [c.id for c in _pack().chapters]
    assert pack_store.latest("databricks").version == LATEST


def test_eighteen_written_scenarios_numbered_in_order():
    scenarios = _scenarios()
    assert [s.number for s in scenarios] == list(range(1, 19))
    assert [s.id for s in scenarios] == [str(n) for n in range(1, 19)]
    assert all(s.content for s in scenarios)


def test_each_scenario_cites_real_notes_sections_and_names_a_real_chapter():
    notes = NOTES.read_text(encoding="utf-8")
    sections = {int(n) for n in re.findall(r"^## (\d+)\. ", notes, re.M)}
    chapters = {c.id for c in _pack().chapters}
    for s in _scenarios():
        cited = {int(n) for n in re.findall(r"§(\d+)", s.sources)}
        assert cited and cited <= sections, f"scenario {s.id} cites {sorted(cited - sections) or 'nothing'}"
        assert s.chapter in chapters, f"scenario {s.id} names chapter {s.chapter!r}"


def test_each_scenario_has_the_shape_the_scenario_page_reads():
    """The page keys lists by their text, so repeated lines in one list would collide."""
    for s in _scenarios():
        c = s.content
        where = f"scenario {s.id}"
        assert c["bookmark"].strip() and c["honesty"].strip(), where
        assert len(c["check"]) == 4, where
        for q in c["check"]:
            assert len(q["options"]) == len(set(q["options"])) >= 3, where
            assert 0 <= q["answer"] < len(q["options"]) and q["why"].strip(), where
        case = c["caseStudy"]
        assert case["setting"].strip() and case["tasks"] and case["pipeline"], where
        assert "(fictional)" in case["setting"], f"{where}: say the organisation is fictional"
        for key_list in ([e["when"] for e in case["events"]], case["pipeline"], case["tasks"], c["takeaway"]):
            assert len(key_list) == len(set(key_list)), where
        assert len(c["debrief"]) >= 2, where
        assert set(c["lenses"]) == ROLES, where
        for role, lens in c["lenses"].items():
            assert lens["tasks"] and len(lens["tasks"]) == len(set(lens["tasks"])), f"{where}/{role}"
            assert lens["debrief"], f"{where}/{role}"
            assert lens["sayIt"]["question"].strip() and len(lens["sayIt"]["points"]) >= 3, f"{where}/{role}"
        for block in c["debrief"] + [b for lens in c["lenses"].values() for b in lens["debrief"]]:
            assert block["title"].strip() and len(block["text"]) == len(set(block["text"])), where


def test_chapters_link_to_the_scenarios_that_practise_them():
    by_chapter = {}
    for s in _scenarios():
        by_chapter.setdefault(s.chapter, []).append(s.id)
    for chapter in _pack().chapters:
        assert [link.scenario_id for link in chapter.practice_links] == by_chapter.get(chapter.id, []), chapter.id


def test_correct_answers_are_spread_over_the_option_positions():
    """If the right answer is nearly always in one place, the check tests guessing, not understanding."""
    from collections import Counter

    positions = Counter(q["answer"] for s in _scenarios() for q in s.content["check"])
    assert max(positions.values()) <= sum(positions.values()) / 2, positions


def test_the_right_answer_is_not_usually_the_longest_option():
    """A check passable by always picking the longest option tests guessing (the banks hold the same line)."""
    checks = [q for s in _scenarios() for q in s.content["check"]]
    longest = sum(len(q["options"][q["answer"]]) == max(len(o) for o in q["options"]) for q in checks)
    assert longest / len(checks) <= 0.40, f"{longest} of {len(checks)}"


def test_scenarios_use_current_names_never_teach_deny_and_state_no_pass_mark():
    for s in _scenarios():
        for sentence in _scenario_sentences(s):
            for name in FORMER_NAMES:
                if re.search(r"\b" + re.escape(name) + r"\b", sentence):
                    assert FORMER_MARKERS.search(sentence), f"scenario {s.id}: {name!r} in {sentence!r}"
            assert not re.search(r"pass(ing)? (mark|score)", sentence, re.IGNORECASE), f"scenario {s.id}: {sentence!r}"
        # A DENY may appear as a wrong option or beside what the docs say about it; never as a right answer.
        for q in s.content["check"]:
            right = q["options"][q["answer"]]
            assert "DENY" not in right or re.search(r"no DENY|legacy|ABAC|Beta", right), f"scenario {s.id}: {right!r}"
        for block in s.content["debrief"] + [b for lens in s.content["lenses"].values() for b in lens["debrief"]]:
            for line in block["text"]:
                for sentence in _sentences(line):
                    if "DENY" in sentence:
                        assert re.search(r"no DENY|legacy|ABAC|Beta|DENY policy", sentence), (
                            f"scenario {s.id}: {sentence!r}"
                        )
